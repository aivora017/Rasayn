"""Shared docx helpers for PharmaCare Pro legal kit (v0.9 lawyer review draft)."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
NAVY = RGBColor(0x1E,0x3A,0x8A); GREY = RGBColor(0x55,0x55,0x55); HDR="E8EEFB"; BRD="CCCCCC"
def make_doc():
    d=Document(); s=d.sections[0]
    s.page_width=Inches(8.5); s.page_height=Inches(11)
    s.left_margin=Inches(1); s.right_margin=Inches(1); s.top_margin=Inches(1); s.bottom_margin=Inches(1)
    n=d.styles["Normal"]; n.font.name="Arial"; n.font.size=Pt(11)
    for lvl,sz in [(1,15),(2,13),(3,11)]:
        st=d.styles[f"Heading {lvl}"]; st.font.name="Arial"; st.font.bold=True; st.font.size=Pt(sz); st.font.color.rgb=NAVY
    return d
def title(d,t):
    par=d.add_paragraph(); par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=par.add_run(t); r.font.name="Arial"; r.font.bold=True; r.font.size=Pt(18); r.font.color.rgb=NAVY
def subtitle(d,t):
    par=d.add_paragraph(); par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=par.add_run(t); r.font.name="Arial"; r.font.italic=True; r.font.size=Pt(10); r.font.color.rgb=GREY
def h1(d,t): d.add_heading(t,1)
def h2(d,t): d.add_heading(t,2)
def p(d,t,italic=False,bold=False):
    par=d.add_paragraph(); r=par.add_run(t); r.font.name="Arial"; r.font.italic=italic; r.font.bold=bold; r.font.size=Pt(11)
def bul(d,t):
    par=d.add_paragraph(style="List Bullet"); par.text=""; r=par.add_run(t); r.font.name="Arial"; r.font.size=Pt(11)
def num(d,t):
    par=d.add_paragraph(style="List Number"); par.text=""; r=par.add_run(t); r.font.name="Arial"; r.font.size=Pt(11)
def sp(d): d.add_paragraph("")
def br(d): d.add_page_break()
def _shade(c,fill):
    tcPr=c._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd")
    shd.set(qn("w:val"),"clear"); shd.set(qn("w:color"),"auto"); shd.set(qn("w:fill"),fill); tcPr.append(shd)
def _bord(c):
    tcPr=c._tc.get_or_add_tcPr(); tcb=OxmlElement("w:tcBorders")
    for e in ("top","left","bottom","right"):
        b=OxmlElement(f"w:{e}"); b.set(qn("w:val"),"single"); b.set(qn("w:sz"),"4"); b.set(qn("w:color"),BRD); tcb.append(b)
    tcPr.append(tcb)
def tbl(d,rows,fill_header=True):
    t=d.add_table(rows=len(rows),cols=len(rows[0])); t.autofit=True
    for ri,row in enumerate(rows):
        for ci,v in enumerate(row):
            c=t.rows[ri].cells[ci]; c.text=""
            par=c.paragraphs[0]; r=par.add_run(str(v)); r.font.name="Arial"; r.font.size=Pt(10); r.font.bold=(ri==0)
            _bord(c)
            if ri==0 and fill_header: _shade(c,HDR)
            c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
    return t
def hf(d,hdr_text,ftr_text):
    s=d.sections[0]
    h=s.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rh=h.add_run(hdr_text); rh.font.name="Arial"; rh.font.size=Pt(9); rh.font.color.rgb=GREY
    f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.LEFT
    rf=f.add_run(ftr_text+"    Page "); rf.font.name="Arial"; rf.font.size=Pt(9); rf.font.color.rgb=GREY
    fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); f._p.append(fld)
    rf2=f.add_run(" of "); rf2.font.name="Arial"; rf2.font.size=Pt(9); rf2.font.color.rgb=GREY
    fld2=OxmlElement("w:fldSimple"); fld2.set(qn("w:instr"),"NUMPAGES"); f._p.append(fld2)

def save_doc(d, out_path):
    """Save document and patch settings.xml to add w:percent="100" to <w:zoom/>."""
    import zipfile, os, re
    d.save(out_path)
    tmp = out_path + ".tmp"
    with zipfile.ZipFile(out_path, 'r') as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item == "word/settings.xml":
                s = data.decode('utf-8')
                s = re.sub(r'<w:zoom(?![^>]*w:percent)([^/]*)/>', r'<w:zoom\1 w:percent="100"/>', s)
                data = s.encode('utf-8')
            zout.writestr(item, data)
    os.replace(tmp, out_path)
